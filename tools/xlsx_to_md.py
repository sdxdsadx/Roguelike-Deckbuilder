"""把策划表 xlsx 导出成 Markdown，方便在 GitHub 上直接阅读。

用法：
    python3 tools/xlsx_to_md.py design/game-design.xlsx design/game-design.md
"""
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def cell_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    return text.replace("|", "\\|").replace("\r\n", "\n").replace("\n", "<br>")


def sheet_to_md(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return f"## {ws.title}\n\n（空表）\n"
    header = rows[0]
    body = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    width = max(len(r) for r in rows)
    # 只保留有表头或有内容的列
    keep = [
        i for i in range(width)
        if (i < len(header) and header[i] not in (None, ""))
        or any(i < len(r) and r[i] not in (None, "") for r in body)
    ]
    titles = [
        cell_text(header[i]) if i < len(header) and header[i] not in (None, "")
        else f"（{get_column_letter(i + 1)} 列）"
        for i in keep
    ]
    lines = [f"## {ws.title}", ""]
    lines.append("| " + " | ".join(titles) + " |")
    lines.append("| " + " | ".join("---" for _ in keep) + " |")
    for r in body:
        lines.append("| " + " | ".join(cell_text(r[i]) if i < len(r) else "" for i in keep) + " |")
    lines.append("")
    return "\n".join(lines)


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    parts = [
        "# 策划数据表",
        "",
        f"> 由 `tools/xlsx_to_md.py` 从 `{src}` 自动生成，请修改 xlsx 后重新生成，不要手改本文件。",
        "",
    ]
    for ws in wb.worksheets:
        parts.append(sheet_to_md(ws))
    with open(dst, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
